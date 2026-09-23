"""
Module 7: Attendance History & Reports.

Provides:
- Searchable attendance history (by name, roll, date range, department, year, section)
- PDF export (via reportlab)
- Excel export (via openpyxl)
"""
from datetime import datetime
from io import BytesIO

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import FileResponse
from sqlalchemy.orm import Session
from sqlalchemy import or_

# PDF export
from reportlab.lib import pagesizes
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.lib import colors

# Excel export
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

from app.database import get_db
from app.models.attendance import Attendance
from app.models.student import Student
from app.models.admin import Admin
from app.core.deps import get_current_admin
from app.schemas.reports import PaginatedAttendanceRecords, AttendanceRecordOut

router = APIRouter(prefix="/api/reports", tags=["Reports"])


def _record_to_out(record: Attendance, student: Student) -> AttendanceRecordOut:
    return AttendanceRecordOut(
        id=record.id,
        student_id=student.id,
        full_name=student.full_name,
        roll_number=student.roll_number,
        department=student.department,
        date=record.date,
        time=record.time,
        status=record.status if isinstance(record.status, str) else record.status.value,
        confidence=record.confidence,
        marked_by=record.marked_by if isinstance(record.marked_by, str) else record.marked_by.value,
    )


@router.get("/attendance-history", response_model=PaginatedAttendanceRecords)
def search_attendance(
    search: str = Query(None, description="Search by name or roll number"),
    department: str = Query(None),
    year: str = Query(None),
    section: str = Query(None),
    date_from: str = Query(None, description="YYYY-MM-DD"),
    date_to: str = Query(None, description="YYYY-MM-DD"),
    page: int = Query(1, ge=1),
    page_size: int = Query(20, ge=1, le=500),
    db: Session = Depends(get_db),
    _admin: Admin = Depends(get_current_admin),
):
    q = db.query(Attendance).join(Student, Attendance.student_id == Student.id)

    if search:
        like = f"%{search.strip()}%"
        q = q.filter(
            or_(Student.full_name.ilike(like), Student.roll_number.ilike(like))
        )
    if department:
        q = q.filter(Student.department == department)
    if year:
        q = q.filter(Student.year == year)
    if section:
        q = q.filter(Student.section == section)

    if date_from:
        try:
            df = datetime.strptime(date_from, "%Y-%m-%d").date()
            q = q.filter(Attendance.date >= df)
        except ValueError:
            pass
    if date_to:
        try:
            dt = datetime.strptime(date_to, "%Y-%m-%d").date()
            q = q.filter(Attendance.date <= dt)
        except ValueError:
            pass

    total = q.count()
    records = (
        q.order_by(Attendance.date.desc(), Attendance.time.desc())
        .offset((page - 1) * page_size)
        .limit(page_size)
        .all()
    )

    items = [_record_to_out(r, r.student) for r in records]
    return PaginatedAttendanceRecords(
        total=total, page=page, page_size=page_size, items=items
    )


@router.get("/export-pdf")
def export_pdf(
    search: str = Query(None),
    department: str = Query(None),
    year: str = Query(None),
    section: str = Query(None),
    date_from: str = Query(None),
    date_to: str = Query(None),
    db: Session = Depends(get_db),
    _admin: Admin = Depends(get_current_admin),
):
    """Export filtered attendance records as PDF."""
    q = db.query(Attendance).join(Student, Attendance.student_id == Student.id)

    filter_desc = []
    if search:
        like = f"%{search.strip()}%"
        q = q.filter(
            or_(Student.full_name.ilike(like), Student.roll_number.ilike(like))
        )
        filter_desc.append(f"Search: {search}")
    if department:
        q = q.filter(Student.department == department)
        filter_desc.append(f"Dept: {department}")
    if year:
        q = q.filter(Student.year == year)
        filter_desc.append(f"Year: {year}")
    if section:
        q = q.filter(Student.section == section)
        filter_desc.append(f"Section: {section}")
    if date_from:
        try:
            df = datetime.strptime(date_from, "%Y-%m-%d").date()
            q = q.filter(Attendance.date >= df)
            filter_desc.append(f"From: {df}")
        except ValueError:
            pass
    if date_to:
        try:
            dt = datetime.strptime(date_to, "%Y-%m-%d").date()
            q = q.filter(Attendance.date <= dt)
            filter_desc.append(f"To: {dt}")
        except ValueError:
            pass

    records = q.order_by(Attendance.date.desc()).all()

    # Build PDF
    pdf_buffer = BytesIO()
    doc = SimpleDocTemplate(pdf_buffer, pagesize=pagesizes.letter)
    story = []

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "CustomTitle",
        parent=styles["Heading1"],
        fontSize=16,
        textColor=colors.HexColor("#1f2937"),
        spaceAfter=6,
    )

    story.append(Paragraph("Attendance Report", title_style))
    story.append(Spacer(1, 0.2 * inch))

    if filter_desc:
        filter_text = " | ".join(filter_desc)
        story.append(Paragraph(f"<b>Filters:</b> {filter_text}", styles["Normal"]))
        story.append(Spacer(1, 0.1 * inch))

    story.append(
        Paragraph(
            f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}",
            styles["Normal"],
        )
    )
    story.append(Spacer(1, 0.3 * inch))

    # Table data
    table_data = [
        ["Date", "Student", "Roll No.", "Department", "Time", "Status", "Marked By"]
    ]
    for record in records:
        table_data.append([
            record.date.isoformat(),
            record.student.full_name,
            record.student.roll_number,
            record.student.department,
            record.time.strftime("%H:%M:%S") if record.time else "",
            record.status.value,
            record.marked_by.value,
        ])

    if table_data:
        table = Table(table_data, colWidths=[1.0*inch, 1.5*inch, 1.0*inch, 1.0*inch, 0.8*inch, 0.8*inch, 0.9*inch])
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#3b6ff2")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.whitesmoke),
            ("ALIGN", (0, 0), (-1, -1), "CENTER"),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, 0), 10),
            ("BOTTOMPADDING", (0, 0), (-1, 0), 12),
            ("BACKGROUND", (0, 1), (-1, -1), colors.beige),
            ("GRID", (0, 0), (-1, -1), 1, colors.black),
        ]))
        story.append(table)

    doc.build(story)
    pdf_buffer.seek(0)

    return FileResponse(
        pdf_buffer,
        media_type="application/pdf",
        filename=f"attendance-report-{datetime.now().strftime('%Y%m%d-%H%M%S')}.pdf",
    )


@router.get("/export-excel")
def export_excel(
    search: str = Query(None),
    department: str = Query(None),
    year: str = Query(None),
    section: str = Query(None),
    date_from: str = Query(None),
    date_to: str = Query(None),
    db: Session = Depends(get_db),
    _admin: Admin = Depends(get_current_admin),
):
    """Export filtered attendance records as Excel."""
    q = db.query(Attendance).join(Student, Attendance.student_id == Student.id)

    filter_desc = []
    if search:
        like = f"%{search.strip()}%"
        q = q.filter(
            or_(Student.full_name.ilike(like), Student.roll_number.ilike(like))
        )
        filter_desc.append(f"Search: {search}")
    if department:
        q = q.filter(Student.department == department)
        filter_desc.append(f"Dept: {department}")
    if year:
        q = q.filter(Student.year == year)
        filter_desc.append(f"Year: {year}")
    if section:
        q = q.filter(Student.section == section)
        filter_desc.append(f"Section: {section}")
    if date_from:
        try:
            df = datetime.strptime(date_from, "%Y-%m-%d").date()
            q = q.filter(Attendance.date >= df)
            filter_desc.append(f"From: {df}")
        except ValueError:
            pass
    if date_to:
        try:
            dt = datetime.strptime(date_to, "%Y-%m-%d").date()
            q = q.filter(Attendance.date <= dt)
            filter_desc.append(f"To: {dt}")
        except ValueError:
            pass

    records = q.order_by(Attendance.date.desc()).all()

    # Build Excel
    wb = Workbook()
    ws = wb.active
    ws.title = "Attendance"

    # Header styling
    header_fill = PatternFill(start_color="3b6ff2", end_color="3b6ff2", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=11)
    border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    headers = ["Date", "Student", "Roll No.", "Department", "Time", "Status", "Marked By"]
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col)
        cell.value = header
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border

    # Set column widths
    ws.column_dimensions["A"].width = 12
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 12
    ws.column_dimensions["D"].width = 15
    ws.column_dimensions["E"].width = 10
    ws.column_dimensions["F"].width = 10
    ws.column_dimensions["G"].width = 10

    # Data rows
    for row, record in enumerate(records, start=2):
        row_data = [
            record.date.isoformat(),
            record.student.full_name,
            record.student.roll_number,
            record.student.department,
            record.time.strftime("%H:%M:%S") if record.time else "",
            record.status.value,
            record.marked_by.value,
        ]
        for col, value in enumerate(row_data, start=1):
            cell = ws.cell(row=row, column=col)
            cell.value = value
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = border

    # Add filter info as a note at the top
    if filter_desc:
        ws.insert_rows(1)
        ws.cell(row=1, column=1).value = f"Filters: {' | '.join(filter_desc)}"
        ws.cell(row=1, column=1).font = Font(bold=True, size=10)

    # Save to BytesIO
    excel_buffer = BytesIO()
    wb.save(excel_buffer)
    excel_buffer.seek(0)

    return FileResponse(
        excel_buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        filename=f"attendance-report-{datetime.now().strftime('%Y%m%d-%H%M%S')}.xlsx",
    )
